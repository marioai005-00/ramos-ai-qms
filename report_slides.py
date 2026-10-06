"""Draws the 8D report on the company slide format (assets/report_template.pptx).

This is the only module that needs python-pptx. It only lays out what the report agent hands over:
the saved Case, the agent's calculations, its check results and the AI summary draft. It adds no content.
"""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from closure_advisor import CAUSES, KST, STAGES, parse_time
from tool_advisor import _dict, _rows, _text

TEMPLATE = Path(__file__).resolve().parent / "assets" / "report_template.pptx"
FONT = "맑은 고딕"
NAVY, INK, MUTED, WHITE = "002060", "1F2937", "4B5563", "FFFFFF"
BLUE, GREEN, PURPLE, RED, AMBER = "1F4E9E", "1E7B45", "6A3FC8", "B42318", "B45309"
TINT = {BLUE: "EAF0FA", GREEN: "E7F4EC", PURPLE: "F0EBFB", RED: "FDECEA", AMBER: "FDF1E7", NAVY: "E9EDF5", MUTED: "F1F3F6"}
LEFT, WIDTH = 0.5, 12.33
LEVELS = {"block": ("확인 필요", RED), "warn": ("주의", AMBER), "info": ("참고", MUTED)}


def clip(value: Any, limit: int) -> str:
    text = _text(value, 4000)
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def day(value: Any) -> str:
    parsed = parse_time(value)
    return parsed.strftime("%Y-%m-%d %H:%M") if parsed else _text(value)


def person(value: Any) -> str:
    """'이성우 팀장_P.Pro (Flash 개발3팀)' → '이성우 팀장'."""
    return _text(value).split("(")[0].split("_")[0].strip()


class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self.body = self.prs.slide_layouts[1]
        ids = self.prs.slides._sldIdLst
        for extra in list(ids)[1:]:  # keep the cover, drop the template's sample slides
            self.prs.part.drop_rel(extra.get(qn("r:id")))
            ids.remove(extra)

    # ---- drawing primitives
    @staticmethod
    def _style(run, size, bold=False, color=INK):
        run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", FONT)

    def text(self, slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        """paras: list of str or (text, size, bold, color, bullet)."""
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap, tf.vertical_anchor = True, anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        for i, p in enumerate(paras):
            given = [p] if isinstance(p, str) else list(p)
            value, size, bold, color, bullet = given + ["", 12, False, INK, False][len(given):]
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.space_after = Pt(5 if bullet else 2)
            ppr = para._p.get_or_add_pPr()
            if bullet:
                ppr.set("marL", str(Inches(0.2)))
                ppr.set("indent", str(-Inches(0.16)))
                ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
                ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))
            else:
                ppr.append(ppr.makeelement(qn("a:buNone"), {}))
            run = para.add_run()
            run.text = value
            self._style(run, size, bold, color)
        return box

    def rect(self, slide, x, y, w, h, fill, radius=0.06):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.adjustments[0] = radius
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def pill(self, slide, x, y, label, color, w=None, size=10):
        shape = self.rect(slide, x, y, w or (0.26 + 0.15 * len(label)), 0.28, color, radius=0.5)
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        run = tf.paragraphs[0].add_run()
        run.text = label
        self._style(run, size, True, WHITE)
        return shape

    def stat(self, slide, x, y, w, label, value, color=BLUE, note="", h=1.05):
        self.rect(slide, x, y, w, h, TINT[color])
        self.text(slide, x + 0.15, y + 0.08, w - 0.3, 0.28, [(label, 10, False, MUTED)])
        self.text(slide, x + 0.15, y + 0.34, w - 0.3, 0.42, [(value, 18 if len(value) < 14 else 13, True, color)])
        if note:
            self.text(slide, x + 0.15, y + h - 0.3, w - 0.3, 0.26, [(note, 9, False, MUTED)])

    def box(self, slide, x, y, w, h, title, lines, color=NAVY, size=11):
        self.rect(slide, x, y, w, h, TINT[color])
        self.text(slide, x + 0.15, y + 0.08, w - 0.3, 0.3, [(title, 12, True, color)])
        self.text(slide, x + 0.15, y + 0.42, w - 0.3, h - 0.5, [(line, size, False, INK, len(lines) > 1) for line in lines])

    def table(self, slide, x, y, w, rows, widths, size=10, row_h=0.36):
        shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
        tbl = shape.table
        for i, share in enumerate(widths):
            tbl.columns[i].width = Inches(w * share)
        for r, row in enumerate(rows):
            tbl.rows[r].height = Inches(row_h)
            for c, value in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else ("FFFFFF" if r % 2 else "F5F7FA"))
                cell.margin_left = cell.margin_right = Inches(0.07)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text_frame.word_wrap = True
                color, bold = (WHITE, True) if r == 0 else (INK, False)
                if isinstance(value, tuple):
                    value, color, bold = value
                run = cell.text_frame.paragraphs[0].add_run()
                run.text = str(value)
                self._style(run, size, bold, color)
        return shape

    def slide(self, title, stage_badge=None):
        slide = self.prs.slides.add_slide(self.body)
        slide.shapes.title.text_frame.text = " " + title
        for run in slide.shapes.title.text_frame.paragraphs[0].runs:
            run.font.name = FONT
        if stage_badge:
            label, color = stage_badge
            self.pill(slide, LEFT + WIDTH - 2.9, 0.78, label, color, w=2.9)
        return slide

    def bytes(self) -> bytes:
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _badge(case: dict, stage: str):
    history = _dict(_dict(case.get("signOffHistory")).get(stage))
    if history.get("status") == "Approved":
        return f"{stage} 결재 완료 · {day(_dict(history.get('champion')).get('signedAt'))}", GREEN
    return f"{stage} 결재 전 — 확정되지 않은 기록", RED


def build(case: dict, calc: dict, checks: list[dict], summary: dict | None, evidence: dict, meta: dict) -> tuple[bytes, int]:
    """Returns the .pptx bytes and the slide count."""
    deck = Deck()
    d2, d3, d4, d5, d6, d7, d8 = (_dict(case.get(k)) for k in ("d2", "d3", "d4", "d5", "d6", "d7", "d8"))
    closed = case.get("status") == "Closed"

    # ---- cover
    cover = deck.prs.slides[0]
    for shape in cover.shapes:
        if shape.is_placeholder:
            shape.text_frame.text = "8D Report"
            for run in shape.text_frame.paragraphs[0].runs:
                run.font.size, run.font.name = Pt(32), FONT
        elif shape.has_text_frame and shape.text_frame.text.strip() == "YYYY. MM. DD":
            shape.text_frame.paragraphs[0].runs[0].text = datetime.now(KST).strftime("%Y. %m. %d")
    deck.text(cover, 2.2, 2.75, 7.8, 0.9, [(clip(case.get("claimTitle"), 60), 18, True, WHITE),
                                           (f"{clip(case.get('customer'), 24)} · {clip(case.get('product'), 40)}", 14, False, WHITE)], align=PP_ALIGN.RIGHT)
    deck.text(cover, 0.45, 5.45, 9.6, 0.9, [(f"Case {case.get('id')} · {'종결' if closed else '진행 중 (' + _text(case.get('currentStage')) + ')'}", 12, True, NAVY),
                                            ("8D 보고서 에이전트가 저장된 기록으로 작성한 초안입니다. 수치·판정·결재는 기록 그대로이며, 송부 전 사람이 확인합니다.", 10, False, MUTED)]
              + ([("검증용 가상 자료로 만든 보고서입니다 (실제 고객 건 아님).", 10, True, RED)] if meta.get("testData") else []))

    # ---- summary
    s = deck.slide("요약")
    facts = [("항목", "기록"), ("고객사", clip(case.get("customer"), 40)), ("제품 / 품번", clip(f"{_text(case.get('product'))} / {_text(case.get('partNumber'))}", 62)),
             ("Lot", clip(case.get("lotNumber"), 50)), ("불량 현상", clip(case.get("claimTitle"), 62)),
             ("불량 / 검사 수량", calc.get("defectLine") or "기록 없음"), ("접수", day(case.get("receiptDate"))),
             ("상태", "종결 · " + day(case.get("closedAt")) if closed else f"진행 중 · 현재 {_text(case.get('currentStage'))}"),
             ("결재 완료 단계", f"{calc['approvedStages']} / 8")]
    deck.table(s, LEFT, 1.2, 6.0, facts, [0.3, 0.7], size=11, row_h=0.4)
    deck.rect(s, 6.75, 1.2, 6.08, 3.6, TINT[PURPLE])
    if summary:
        deck.pill(s, 6.9, 1.32, "AI 초안 · 사람 확인 필요", PURPLE, w=2.3)
        deck.text(s, 6.9, 1.7, 5.8, 0.6, [(clip(summary.get("headline"), 110), 13, True, INK)])
        deck.text(s, 6.9, 2.35, 5.8, 2.4, [(clip(p, 150), 11, False, INK, True) for p in summary.get("points", [])[:5]])
    else:
        deck.text(s, 6.9, 1.32, 5.8, 0.4, [("요약 초안 없음", 13, True, PURPLE)])
        deck.text(s, 6.9, 1.8, 5.8, 1.0, [("외부 AI 요약을 받지 못했거나 계획에서 제외되어, 요약 문장을 넣지 않았습니다. 단계별 기록은 다음 장부터 그대로 실었습니다.", 11, False, INK)])
    for i, row in enumerate(calc.get("leadTimes", [])[:3]):
        color = {"on_time": GREEN, "late": RED, "overdue": RED}.get(row["status"], MUTED)
        deck.stat(s, LEFT + i * 3.13, 5.0, 2.98, f"{row['label']} · 기한 {row['dueAt'] or '-'}", row["text"], color, note=f"결재 {row['doneAt']}" if row["doneAt"] else "", h=1.2)
    open_checks = [c for c in checks if c["level"] in {"block", "warn"}]
    deck.stat(s, LEFT + 3 * 3.13, 5.0, 2.94, "에이전트 점검", f"확인할 것 {len(open_checks)}건", RED if any(c["level"] == "block" for c in checks) else (AMBER if open_checks else GREEN),
              note="마지막 장에 목록", h=1.2)
    deck.text(s, LEFT, 6.35, WIDTH, 0.3, [("처리 기간은 접수 시각부터 해당 단계 챔피언 결재 시각까지입니다.", 9, False, MUTED)])

    # ---- D1 · D2
    s = deck.slide("D1 팀 구성 · D2 문제 정의", _badge(case, "D2"))
    team = [("역할", "이름", "부서")] + [(clip(m.get("role"), 30), clip(m.get("name"), 16), clip(m.get("dept"), 14)) for m in _rows(case.get("team"))[:10]]
    deck.table(s, LEFT, 1.25, 5.6, team if len(team) > 1 else team + [("기록 없음", "", "")], [0.5, 0.27, 0.23], size=10, row_h=0.42)
    w2h = [("5W2H", "기록")] + [(k, clip(d2.get(f), 78)) for k, f in (("무엇을", "problemWhat"), ("어디서", "problemWhere"), ("언제", "problemWhen"), ("누가", "problemWho"),
                                                                     ("어느 것", "problemWhich"), ("어떻게", "problemHow"), ("얼마나", "problemHowMany"))]
    deck.table(s, 6.35, 1.25, 6.48, w2h, [0.17, 0.83], size=10, row_h=0.5)
    deck.text(s, 6.35, 5.4, 6.48, 0.3, [("문제 정의문", 11, True, NAVY)])
    deck.text(s, 6.35, 5.7, 6.48, 1.1, [(clip(d2.get("problemStatement"), 300) or "기록 없음", 10, False, INK)])

    # ---- D2 IS / IS NOT
    rows = _rows(d2.get("isIsNot"))[:7]
    if rows:
        s = deck.slide("D2 IS / IS NOT — 어디에서만 생겼는가", _badge(case, "D2"))
        body = [("구분", "IS (생긴 곳)", "IS NOT (생기지 않은 곳)", "차이", "사실 확인")]
        body += [(clip(r.get("factor"), 14), clip(r.get("is"), 70), clip(r.get("isNot"), 80), clip(r.get("difference"), 60),
                  ("확인", GREEN, True) if r.get("verificationStatus") == "Verified" else ("미확인", RED, True)) for r in rows]
        deck.table(s, LEFT, 1.25, WIDTH, body, [0.11, 0.26, 0.3, 0.23, 0.1], size=11, row_h=0.62)

    # ---- D3
    s = deck.slide("D3 봉쇄 조치", _badge(case, "D3"))
    for i, (label, value) in enumerate(calc.get("containment", [])[:4]):
        deck.stat(s, LEFT + i * 3.13, 1.2, 2.98 if i < 3 else 2.94, label, value, BLUE)
    actions = [("조치", "대상", "담당", "상태", "결과·근거")] + [(clip(a.get("action"), 44), clip(a.get("target"), 20), person(a.get("owner")), clip(a.get("status"), 12), clip(a.get("result"), 56))
                                                        for a in _rows(d3.get("actions"))[:6]]
    deck.table(s, LEFT, 2.45, WIDTH, actions if len(actions) > 1 else actions + [("기록 없음", "", "", "", "")], [0.3, 0.14, 0.12, 0.09, 0.35], size=10, row_h=0.45)
    deck.box(s, LEFT, 5.45, WIDTH, 1.25, "봉쇄 유효성 (기록)", [clip(d3.get("effectivenessStatement"), 330) or "기록 없음"], BLUE, size=10)

    # ---- D4
    s = deck.slide("D4 근본 원인", _badge(case, "D4"))
    causes = _dict(d4.get("rootCauses"))
    for i, (key, label) in enumerate(CAUSES.items()):
        cause, x, w = _dict(causes.get(key)), LEFT + i * 4.18, 3.98
        confirmed = cause.get("status") == "Confirmed"
        deck.rect(s, x, 1.25, w, 5.45, TINT[BLUE if confirmed else MUTED])
        deck.text(s, x + 0.2, 1.38, 2.2, 0.35, [(label, 15, True, BLUE)])
        deck.pill(s, x + w - 1.35, 1.42, "확정" if confirmed else (clip(cause.get("status"), 8) or "기록 없음"), GREEN if confirmed else MUTED, w=1.15)
        deck.text(s, x + 0.2, 1.9, w - 0.4, 1.5, [(clip(cause.get("statement"), 150) or "기록 없음", 14, True, INK)])
        deck.text(s, x + 0.2, 3.55, w - 0.4, 0.3, [("검증 방법", 11, True, MUTED)])
        deck.text(s, x + 0.2, 3.87, w - 0.4, 1.1, [(clip(cause.get("validationMethod"), 130) or "-", 12, False, INK)])
        deck.text(s, x + 0.2, 5.05, w - 0.4, 0.3, [("근거 파일", 11, True, MUTED)])
        deck.text(s, x + 0.2, 5.37, w - 0.4, 1.25, [(clip(cause.get("evidence"), 150) or "-", 10, False, INK)])

    # ---- D5
    s = deck.slide("D5 영구 대책", _badge(case, "D5"))
    candidates = _rows(d5.get("candidates"))
    selected = [c for c in candidates if c.get("selected") is True]
    body = [("원인", "선정 대책", "원인 제거 근거", "담당", "기한")] + [(CAUSES.get(c.get("causeType"), "미연결"), clip(c.get("title"), 62), clip(c.get("rootCauseElimination"), 54), person(c.get("owner")), clip(c.get("due"), 10))
                                                           for c in selected[:7]]
    deck.table(s, LEFT, 1.25, WIDTH, body if len(body) > 1 else body + [("선정된 대책 없음", "", "", "", "")], [0.1, 0.38, 0.3, 0.12, 0.1], size=12, row_h=0.6)
    pcn = _dict(d5.get("pcnEcn"))
    deck.box(s, LEFT, 5.5, 6.05, 1.2, "검토한 후보", [f"후보 {len(candidates)}건 중 {len(selected)}건 선정 · 선정하지 않은 후보 {len(candidates) - len(selected)}건"], NAVY)
    deck.box(s, 6.78, 5.5, 6.05, 1.2, "고객 변경 승인 (PCN)", [f"PCN 필요: {'예' if pcn.get('pcnRequired') is True else '아니오' if pcn.get('pcnRequired') is False else '기록 없음'} · 승인 상태: {clip(pcn.get('customerApprovalStatus'), 44) or '기록 없음'}"
                                                         + (f" · {clip(pcn.get('ecnNumber'), 24)}" if pcn.get("ecnNumber") else "")], NAVY)

    # ---- D6
    s = deck.slide("D6 효과 검증", _badge(case, "D6"))
    tests = _rows(d6.get("validationTests"))
    body = [("시험", "조건", "합격 기준", "시료 / 불량", "판정", "완료일", "성적서")]
    body += [(clip(t.get("testName"), 22), clip(t.get("condition"), 40), clip(t.get("acceptanceCriteria"), 36), f"{_text(t.get('sampleSize'))} / {_text(t.get('failQty'))}",
              (clip(t.get("result"), 8) or "-", GREEN if t.get("result") == "PASS" else RED, True), clip(t.get("completedAt"), 10), clip(t.get("evidence"), 40)) for t in tests[:6]]
    deck.table(s, LEFT, 1.25, WIDTH, body if len(body) > 1 else body + [("검증 시험 없음", "", "", "", "", "", "")], [0.14, 0.2, 0.18, 0.09, 0.07, 0.09, 0.23], size=10, row_h=0.5)
    stats = calc.get("beforeAfter")
    if stats:
        cards = [("적용 전", f"{stats['before']['ppm']:,} PPM", f"{stats['before']['fail']} / {stats['before']['n']:,}", MUTED),
                 ("적용 후 (관측)", f"{stats['after']['ppm']:,} PPM", f"{stats['after']['fail']} / {stats['after']['n']:,}", BLUE),
                 (f"적용 후 상한 · 신뢰도 {stats['confidence']:.0%}", f"{stats['after']['upperBoundPpm']:,} PPM", f"기준 {stats['target']['ppm']:,} PPM", BLUE),
                 ("판정 (공식)", stats["verdictLabel"], "기록과 재계산 일치" if calc.get("beforeAfterMatches") else "기록과 재계산이 다름 — 확인 필요", GREEN if stats["verdict"] == "improved" else RED)]
        for i, (label, value, note, color) in enumerate(cards):
            deck.stat(s, LEFT + i * 3.13, 4.5, 2.98 if i < 3 else 2.94, label, value, color, note=note, h=1.15)
    else:
        deck.box(s, LEFT, 4.5, WIDTH, 1.15, "개선 전·후 비교", ["전·후 비교 계산 기록이 없습니다."], AMBER)
    release, impl = _dict(d6.get("containmentRelease")), _dict(d6.get("implementationDetails"))
    decision = {"Released": "봉쇄 해제", "Retained": "봉쇄 유지"}.get(release.get("decision"), "결정 기록 없음")
    deck.text(s, LEFT, 5.85, WIDTH, 0.9, [(f"대책 적용: {clip(impl.get('appliedLot'), 30) or '-'} · {clip(impl.get('startDate'), 12) or '-'} · {clip(impl.get('productionSite'), 30) or '-'}", 10, False, INK),
                                          (f"봉쇄 결정: {decision} — {clip(release.get('rationale'), 120) or '사유 기록 없음'}", 10, True, NAVY)])

    # ---- D7
    s = deck.slide("D7 재발 방지", _badge(case, "D7"))
    updates = [("개정 문서", "문서 번호", "개정", "담당", "상태")] + [(clip(u.get("docName"), 26), clip(u.get("docNo"), 24), clip(u.get("rev"), 8), person(u.get("owner")),
                                                         (clip(u.get("status"), 12), GREEN if u.get("status") == "Completed" else RED, True)) for u in _rows(d7.get("systemUpdates"))[:7]]
    deck.table(s, LEFT, 1.25, 7.2, updates if len(updates) > 1 else updates + [("문서 개정 기록 없음", "", "", "", "")], [0.3, 0.3, 0.1, 0.15, 0.15], size=10, row_h=0.45)
    deploy = _rows(d7.get("horizontalDeployment"))
    done = sum(1 for r in deploy if r.get("status") == "Completed")
    na = sum(1 for r in deploy if r.get("status") == "Not Applicable")
    deck.box(s, 7.95, 1.25, 4.88, 2.9, f"수평 전개 {len(deploy)}건", [f"완료 {done}건 · 해당 없음 {na}건 · 미완료 {len(deploy) - done - na}건"] + [clip(r.get("product"), 30) for r in deploy[:6]], BLUE, size=10)
    lesson = _dict(d7.get("lessonsLearned"))
    deck.box(s, LEFT, 4.75, WIDTH, 1.95, "교훈 (다음 Case의 대책 추천에 쓰임)", [clip(lesson.get("lesson"), 260) or "교훈이 저장되지 않았습니다."]
             + ([f"원인: {clip(lesson.get('cause'), 170)}"] if lesson.get("cause") else []), GREEN, size=10)

    # ---- D8
    s = deck.slide("D8 종결 · 결재 이력", _badge(case, "D8"))
    closure = _dict(d8.get("closure"))
    deck.box(s, LEFT, 1.25, 5.6, 2.2, "잔여 위험과 관리", [clip(closure.get("remainingRisk"), 330) or "기록 없음"], AMBER, size=10)
    deck.box(s, LEFT, 3.6, 5.6, 1.35, "고객 수락", [clip(closure.get("customerAcceptance"), 180) or "기록 없음"], NAVY, size=10)
    monitor = _rows(_dict(case.get("postClosureMonitoring")).get("items"))
    deck.box(s, LEFT, 5.1, 5.6, 1.6, "종결 후 재발 확인", [f"{m.get('days')}일 · {clip(m.get('dueDate'), 10)} · {'완료' if m.get('status') != 'Open' else '예정'}" for m in monitor[:4]] or ["계획 기록 없음"], GREEN, size=10)
    history = _dict(case.get("signOffHistory"))
    signs = [("단계", "기안", "리더", "챔피언", "최종 결재")]
    for stage in STAGES:
        h = _dict(history.get(stage))
        name = lambda role: clip(_dict(h.get(role)).get("name"), 8) or "-"
        signs.append((stage, name("drafter"), name("leader"), name("champion"), (day(_dict(h.get("champion")).get("signedAt")), GREEN, False) if h.get("status") == "Approved" else ("결재 전", RED, True)))
    deck.table(s, 6.35, 1.25, 6.48, signs, [0.12, 0.2, 0.2, 0.2, 0.28], size=10, row_h=0.42)
    deck.text(s, 6.35, 5.15, 6.48, 0.6, [("결재는 실제 로그인 계정으로 서버에 기록된 것만 옮겼습니다. 고객 보고서 송부는 종결 조건이 아니며 별도 기록입니다.", 9, False, MUTED)])

    # ---- evidence
    s = deck.slide("근거 자료")
    per_stage = evidence.get("perStage", {})
    body = [("단계", "파일 수", "파일 (일부)")] + [(stage, str(len(per_stage.get(stage, []))), clip(" · ".join(per_stage.get(stage, [])[:4]), 150) or "연결된 근거 없음") for stage in STAGES]
    deck.table(s, LEFT, 1.2, WIDTH, body, [0.07, 0.08, 0.85], size=10, row_h=0.38)
    deck.box(s, LEFT, 4.75, WIDTH, 1.95, "에이전트가 원본을 열어 확인한 것", (evidence.get("lines") or ["근거 원본 열기가 이번 계획에 없었습니다."])[:5], NAVY, size=10)

    # ---- checks
    s = deck.slide("에이전트 점검 결과 — 사람이 확인할 것")
    shown = sorted(checks, key=lambda c: ("block", "warn", "info").index(c["level"]))[:9]
    body = [("구분", "출처", "단계", "내용")]
    for c in shown:
        label, color = LEVELS[c["level"]]
        body.append(((label, color, True), (c["source"], PURPLE if c["source"] == "AI" else INK, c["source"] == "AI"), c.get("stage") or "-", clip(c["text"], 150)))
    if len(body) == 1:
        body.append((("없음", GREEN, True), "-", "-", "규칙 점검과 AI 검토에서 확인할 항목이 나오지 않았습니다."))
    deck.table(s, LEFT, 1.25, WIDTH, body, [0.09, 0.07, 0.07, 0.77], size=10, row_h=0.44)
    if len(checks) > len(shown):
        deck.text(s, LEFT, 5.75, WIDTH, 0.3, [(f"그 밖에 {len(checks) - len(shown)}건은 시스템의 에이전트 실행 기록에 있습니다.", 9, False, MUTED)])
    deck.rect(s, LEFT, 6.1, WIDTH, 0.62, "EEF0F4")
    deck.text(s, LEFT + 0.15, 6.14, WIDTH - 0.3, 0.55, [(f"작성: 8D 보고서 에이전트 · 요청자 {meta.get('requestedBy', '-')} · {datetime.now(KST).strftime('%Y-%m-%d %H:%M')} · 에이전트는 값·판정·결재를 만들지 않았고, "
                                                         "보라색 'AI' 표시는 외부 AI가 쓴 문장입니다.", 10, True, NAVY)], anchor=MSO_ANCHOR.MIDDLE)
    return deck.bytes(), len(deck.prs.slides)
